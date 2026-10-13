from datetime import date, datetime, timezone
from email import policy
from email.parser import BytesParser
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from pharma_suite.automation.alerts import acknowledge, evaluate_alerts
from pharma_suite.automation.config import AutomationConfig
from pharma_suite.automation.reporting import report_path
from pharma_suite.automation.scheduler import build_scheduler, schedule_key
from pharma_suite.automation.schema import alerts, emails, events, executions, reports
from pharma_suite.automation.workflow import enqueue, execute_inline
from pharma_suite.db import create_db, ingest


@pytest.fixture
def automation_engine(tmp_path, hand_data):
    engine = create_db(f"sqlite:///{tmp_path}/auto.db")
    ingest(engine, hand_data)
    return engine


@pytest.fixture
def automation_config(tmp_path):
    return AutomationConfig(
        report_dir=str(tmp_path / "reports"),
        low_stock_units=500,
        assessment_at="2026-01-01T12:00:00+00:00",
    )


def test_config_and_schedule_contract():
    with pytest.raises(ValueError):
        AutomationConfig(daily_cron="broken")
    with pytest.raises(ValueError):
        AutomationConfig(timezone="Unknown/Zone")
    config = AutomationConfig()
    scheduler = build_scheduler(config, lambda kind: None)
    assert {job.id for job in scheduler.get_jobs()} == {"daily", "weekly"}
    instant = datetime(2026, 10, 11, 2, 30, tzinfo=timezone.utc)
    assert schedule_key("daily", instant, config) == schedule_key("daily", instant, config)
    assert schedule_key("weekly", instant, config) != schedule_key("daily", instant, config)


def test_workflow_order_excel_values_and_preview(automation_engine, automation_config):
    run = enqueue(
        automation_engine,
        "daily",
        "hand-1",
        "reviewer",
        "hand fixture",
        automation_config,
        start=date(2025, 1, 1),
        end=date(2025, 12, 31),
    )
    execute_inline(automation_engine, run, automation_config)
    with automation_engine.connect() as conn:
        execution = conn.execute(select(executions)).mappings().one()
        assert execution["status"] == "succeeded"
        row = conn.execute(select(reports)).mappings().one()
        email = conn.execute(select(emails)).mappings().one()
        steps = (
            conn.execute(
                select(events.c.stage)
                .where(events.c.event == "stage_completed")
                .order_by(events.c.id)
            )
            .scalars()
            .all()
        )
        assert steps == ["validation", "analytics", "report", "alerts", "email"]
    book = load_workbook(
        report_path(automation_engine, row["id"], automation_config), data_only=True
    )
    assert book.sheetnames == [
        "Executive Summary",
        "Sales Analytics",
        "Regional Sales",
        "Inventory Levels",
        "Near Expiry",
        "Data Quality",
        "Notes",
    ]
    summary = dict(book["Executive Summary"].iter_rows(min_row=5, max_col=2, values_only=True))
    assert summary["Net sales revenue (INR)"] == 150
    assert summary["Closing stock units"] == 90
    assert summary["Closing stock at cost (INR)"] == 180
    assert book["Regional Sales"]["B5"].value == 150
    assert len(book["Inventory Levels"].tables) == 1
    assert book["Inventory Levels"].freeze_panes == "A5"
    assert email["status"] == "preview"
    message = BytesParser(policy=policy.default).parsebytes(
        Path(automation_config.report_dir, email["filename"]).read_bytes()
    )
    assert "SYNTHETIC" in str(message["Subject"])
    assert any(
        part.get_filename() and part.get_filename().endswith(".xlsx") for part in message.walk()
    )


def test_idempotent_request_and_safe_download(automation_engine, automation_config):
    first = enqueue(automation_engine, "weekly", "unique", "reviewer", "test", automation_config)
    assert (
        enqueue(automation_engine, "weekly", "unique", "reviewer", "test", automation_config)
        == first
    )
    with pytest.raises(ValueError, match="different"):
        enqueue(automation_engine, "daily", "unique", "reviewer", "test", automation_config)
    execute_inline(automation_engine, first, automation_config)
    execute_inline(automation_engine, first, automation_config)
    with automation_engine.connect() as conn:
        assert conn.scalar(select(func.count()).select_from(reports)) == 1
        assert conn.scalar(select(func.count()).select_from(emails)) == 1
    with pytest.raises(ValueError):
        report_path(automation_engine, "../secret", automation_config)


def test_alert_lifecycle_deduplicates_unchanged(automation_engine, automation_config):
    first = evaluate_alerts(automation_engine, automation_config)
    second = evaluate_alerts(automation_engine, automation_config)
    assert first and not second
    with automation_engine.connect() as conn:
        alert = conn.execute(select(alerts).where(alerts.c.kind == "low_stock")).mappings().first()
        count = conn.scalar(select(func.count()).select_from(alerts))
    acknowledge(automation_engine, alert["id"], "reviewer", "Investigating synthetic stock")
    evaluate_alerts(automation_engine, automation_config)
    with automation_engine.connect() as conn:
        assert (
            conn.scalar(select(alerts.c.status).where(alerts.c.id == alert["id"])) == "acknowledged"
        )
        assert conn.scalar(select(func.count()).select_from(alerts)) == count
    quiet = automation_config.model_copy(update={"low_stock_units": 0, "near_expiry_days": 0})
    evaluate_alerts(automation_engine, quiet)
    with automation_engine.connect() as conn:
        assert conn.scalar(select(alerts.c.status).where(alerts.c.id == alert["id"])) == "resolved"
    evaluate_alerts(automation_engine, automation_config)
    with automation_engine.connect() as conn:
        assert conn.scalar(select(alerts.c.status).where(alerts.c.id == alert["id"])) == "new"


def test_failed_workflow_creates_failure_alert(automation_engine, automation_config, monkeypatch):
    import pharma_suite.automation.workflow as workflow

    def fail(*args, **kwargs):
        raise ValueError("Synthetic invalid source")

    monkeypatch.setattr(workflow, "snapshot", fail)
    run = enqueue(
        automation_engine, "daily", "failure", "reviewer", "failure test", automation_config
    )
    execute_inline(automation_engine, run, automation_config)
    with automation_engine.connect() as conn:
        assert conn.scalar(select(executions.c.status)) == "failed"
        assert conn.scalar(select(func.count()).select_from(reports)) == 0
        assert (
            conn.scalar(
                select(func.count()).select_from(alerts).where(alerts.c.kind == "workflow_failure")
            )
            == 1
        )


def test_mail_preview_external_off_by_default(automation_engine, automation_config, monkeypatch):
    from pharma_suite.automation.mail import deliver

    monkeypatch.setenv("SMTP_HOST", "smtp.example.invalid")
    monkeypatch.setenv("SMTP_RECIPIENTS", "recipient@example.invalid")
    run = enqueue(
        automation_engine, "daily", "mail-safe", "reviewer", "safe preview", automation_config
    )
    execute_inline(automation_engine, run, automation_config)
    with automation_engine.connect() as conn:
        mail_id = conn.scalar(select(emails.c.id))
    assert deliver(automation_engine, mail_id, automation_config) == "preview"
