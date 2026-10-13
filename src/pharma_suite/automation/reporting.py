"""XLSX reports from a consistent database snapshot and validated KPI calculations."""

import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select

from pharma_suite.analytics import build_report, facts_from_connection
from pharma_suite.automation.schema import reports
from pharma_suite.automation.storage import atomic_write, digest, safe_artifact
from pharma_suite.generate import TABLE_ORDER
from pharma_suite.quality.demo import make_bundle
from pharma_suite.quality.store import fingerprint, lock, now
from pharma_suite.schema import tables

SUMMARY = {
    "revenue": "Net sales revenue (INR)",
    "units": "Units distributed",
    "invoices": "Distinct invoices",
    "stock_units": "Closing stock units",
    "stock_value": "Closing stock at cost (INR)",
    "risk_value": "Stock at risk (INR)",
    "turnover": "Inventory turnover",
    "near_expiry_units": "Near expiry units",
    "expired_units": "Expired units held",
}


def snapshot(engine):
    with engine.connect() as conn:
        if conn.dialect.name == "postgresql":
            conn = conn.execution_options(isolation_level="REPEATABLE READ")
        with conn.begin():
            lock(conn)
            data = {
                name: [
                    dict(row)
                    for row in conn.execute(
                        select(tables[name]).order_by(*tables[name].primary_key.columns)
                    ).mappings()
                ]
                for name in TABLE_ORDER
            }
            facts = facts_from_connection(conn)
    if not data["inventory"] or not data["sales"]:
        raise ValueError("Business dataset is empty")
    content_hash = fingerprint(data)
    bundle = make_bundle(data, "production-" + content_hash[:24])
    cutoff = min(row["snapshot_date"] for row in data["inventory"])
    bundle.snapshot_date = cutoff.isoformat()
    bundle.feed_date = max(row["sale_date"] for row in data["sales"]).isoformat()
    bundle.reference_at = datetime.combine(
        cutoff + timedelta(days=1), datetime.min.time(), timezone.utc
    )
    # No ingestion timestamp exists in Phase 1; extraction is not a new source import.
    bundle.source_received_at = None
    return bundle, facts, content_hash, cutoff


def calculate(facts, kind, start, end, cutoff, config):
    end = date.fromisoformat(end) if end else cutoff
    start = (
        date.fromisoformat(start) if start else end - timedelta(days=0 if kind == "daily" else 6)
    )
    report = build_report(*facts, start, end, risk_days=config.near_expiry_days)
    return report, start, end


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "item"):
        return value.item()
    return value


def sheet(book, title, headers, rows, note):
    page = book.create_sheet(title)
    page.append([f"NIVARA / {title.upper()} / SYNTHETIC"])
    page.append([note])
    page.append([])
    page.append(headers)
    for row in rows:
        page.append([safe_cell(value) for value in row])
    page.freeze_panes = "A5"
    page.sheet_view.showGridLines = False
    page.row_dimensions[1].height = 30
    for cell in page[1]:
        cell.font = Font(name="Georgia", size=16, color="255F50", bold=True)
    page["A2"].font = Font(size=10, color="666666", italic=True)
    for cell in page[4]:
        cell.fill = PatternFill("solid", fgColor="255F50")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for index, header in enumerate(headers, 1):
        letter = get_column_letter(index)
        page.column_dimensions[letter].width = min(52, max(17, len(header) + 3))
        for row in range(5, page.max_row + 1):
            cell = page.cell(row, index)
            cell.font = Font(name="Calibri", size=11)
            if isinstance(cell.value, (int, float)):
                cell.number_format = (
                    "#,##0.00"
                    if any(
                        x in header.lower() for x in ("revenue", "value", "cost", "price", "inr")
                    )
                    else "#,##0"
                )
    if page.max_row >= 5:
        table = Table(
            displayName="T" + str(len(book.sheetnames)),
            ref=f"A4:{get_column_letter(len(headers))}{page.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        page.add_table(table)
    return page


def write_report(engine, run_id, config, report, start, end, snapshot_hash, quality):
    with engine.connect() as conn:
        existing = (
            conn.execute(select(reports).where(reports.c.execution_id == run_id)).mappings().first()
        )
    if existing:
        if existing["snapshot_hash"] != snapshot_hash:
            raise ValueError("Retry snapshot differs from the registered report snapshot")
        safe_artifact(config, existing["filename"], existing["sha256"])
        return dict(existing)
    report_id = str(uuid.uuid5(uuid.NAMESPACE_URL, "nivara-report:" + run_id))
    filename = f"synthetic_{report_id}.xlsx"
    from pathlib import Path

    path = Path(config.report_dir) / filename
    book = Workbook()
    book.remove(book.active)
    summary = {key: str(report[key]) if report[key] is not None else None for key in SUMMARY}
    summary.update(
        start=start.isoformat(),
        end=end.isoformat(),
        quality_score=quality.overall_score,
        quality_run_id=quality.run_id,
        quality_findings=len(quality.findings),
    )
    executive = sheet(
        book,
        "Executive Summary",
        ["KPI", "Value"],
        [[label, report[key]] for key, label in SUMMARY.items()]
        + [
            ["Data quality score (%)", quality.overall_score],
            ["Quality findings", len(quality.findings)],
        ],
        f"{start} to {end} inclusive | INR, ex-tax | report {report_id}",
    )
    for row in (5, 9, 10):
        executive.cell(row, 2).number_format = "#,##0.00"
    executive["B11"].number_format = '0.0000"x"'
    sheet(
        book,
        "Sales Analytics",
        list(report["monthly"].columns),
        report["monthly"].itertuples(index=False, name=None),
        "Monthly net sales; periods follow the selected window. Revenue = units × selling price.",
    )
    sheet(
        book,
        "Regional Sales",
        list(report["regions"].columns),
        report["regions"].itertuples(index=False, name=None),
        "Attributed to fulfilling warehouse; invoice/line joins do not multiply batch stock.",
    )
    cols = [
        "batch_id",
        "batch_number",
        "medicine",
        "region",
        "warehouse",
        "expiry_date",
        "stock_units",
        "stock_value",
        "days_to_expiry",
        "status",
    ]
    stock = sheet(
        book,
        "Inventory Levels",
        cols,
        report["stock"][cols].itertuples(index=False, name=None),
        f"Closing stock at {end}; receipts less all historical sales, valued at acquisition cost.",
    )
    if stock.max_row >= 5:
        stock.conditional_formatting.add(
            f"G5:G{stock.max_row}",
            CellIsRule(
                operator="lessThan",
                formula=[str(config.low_stock_units)],
                fill=PatternFill("solid", fgColor="FFF0C2"),
            ),
        )
        stock.conditional_formatting.add(
            f"I5:I{stock.max_row}",
            ColorScaleRule(
                start_type="min", start_color="E9AD7A", end_type="max", end_color="DCE8D9"
            ),
        )
    near = report["risk"][report["risk"].days_to_expiry >= 0]
    sheet(
        book,
        "Near Expiry",
        cols,
        near[cols].itertuples(index=False, name=None),
        f"Positive held stock expiring from {end} through {report['risk_cutoff']}. "
        "Expired held stock is shown in Inventory Levels.",
    )
    qcols = ["rule_id", "severity", "entity", "record_ref", "field", "description", "remediation"]
    sheet(
        book,
        "Data Quality",
        qcols,
        [[getattr(f, c) for c in qcols] for f in quality.findings],
        f"Quality run {quality.run_id}, assessment {quality.reference_at.isoformat()}; "
        "derived line/invoice totals are calculated, not independent source declarations.",
    )
    notes = [
        [
            "Data classification",
            "100% synthetic; no patient data. Not a validated GxP or regulatory system.",
        ],
        ["Workflow", run_id],
        ["Database snapshot hash", snapshot_hash],
        ["Quality source", quality.dataset],
        ["Quality assessment clock", quality.reference_at.isoformat()],
        [
            "Timeliness limitation",
            "Source import timestamp unavailable in Phase 1. It is reported "
            "missing, never fabricated from extraction time.",
        ],
        [
            "Inventory",
            "Expired stock legitimately remains held at book cost when "
            "unavailable. No transfers, returns or disposal movements.",
        ],
        ["Turnover", "Period COGS / exact daily average end-of-day inventory. Not annualised."],
        [
            "Low stock alerts",
            "Aggregate medicine/warehouse snapshot units; workbook "
            "highlighting is batch-level review assistance.",
        ],
        [
            "Operational alert clock",
            "Expiry alerts use stored stock snapshot dates, not a claim of "
            "current physical availability.",
        ],
        [
            "Production approval",
            "Reports with noncritical findings remain review artifacts; "
            "workflow never approves or repairs source data.",
        ],
    ]
    sheet(
        book,
        "Notes",
        ["Topic", "Explanation"],
        notes,
        "Read these assumptions before interpreting results.",
    )
    atomic_write(path, lambda temporary: book.save(temporary))
    value = {
        "id": report_id,
        "execution_id": run_id,
        "filename": filename,
        "sha256": digest(path),
        "byte_size": path.stat().st_size,
        "snapshot_hash": snapshot_hash,
        "summary_json": summary,
        "created_at": now(),
    }
    with engine.begin() as conn:
        conn.execute(reports.insert().values(**value))
    return value


def report_path(engine, report_id, config):
    try:
        uuid.UUID(report_id)
    except ValueError as exc:
        raise ValueError("Invalid report identifier") from exc
    with engine.connect() as conn:
        row = conn.execute(select(reports).where(reports.c.id == report_id)).mappings().first()
    if not row:
        raise ValueError("Unknown report")
    return safe_artifact(config, row["filename"], row["sha256"])
